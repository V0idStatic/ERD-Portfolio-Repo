import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SQL_PATH = ROOT / "collieAI.sql"
SCHEMA_OUTPUT = ROOT / "CollieAI_Database_Schema_Copy_Ready.docx"
REL_OUTPUT = ROOT / "CollieAI_Table_Relationships_Copy_Ready.docx"


TABLE_DESCRIPTIONS = {
    "ai_messages": "Stores messages exchanged during AI tutoring sessions.",
    "ai_sessions": "Stores AI tutoring sessions for students and math problems.",
    "api_request_logs": "Records API requests for monitoring and troubleshooting.",
    "attempt_steps": "Stores individual solution steps within a student attempt.",
    "avatar_items": "Defines avatar items that students can unlock.",
    "avatars": "Defines available student avatars.",
    "billing_accounts": "Stores billing identities for individual users or organizations.",
    "class_sections": "Stores classroom sections managed by teachers and organizations.",
    "dda_events": "Records dynamic difficulty adjustment events.",
    "difficulty_levels": "Defines difficulty levels used by math problems and recommendations.",
    "file_assets": "Stores metadata for uploaded and generated files.",
    "guardrail_checks": "Records safety and prompt-injection checks for AI messages.",
    "learning_recommendations": "Stores personalized learning recommendations.",
    "math_problems": "Stores mathematics problems, answers, and solution patterns.",
    "math_skills": "Stores mathematics skills grouped under topics.",
    "math_topics": "Stores high-level mathematics topics.",
    "ml_models": "Stores machine-learning model versions and evaluation metrics.",
    "n8n_workflow_logs": "Records n8n workflow executions associated with sessions.",
    "notifications": "Stores notifications sent to users about students or system events.",
    "ocr_logs": "Records optical character recognition processing results.",
    "organization_members": "Associates teachers with schools or enterprise organizations.",
    "organization_student_enrollments": "Associates students with organizations independently of classes.",
    "organizations": "Stores school and enterprise organization records.",
    "parent_profiles": "Stores parent and guardian profile information.",
    "payment_transactions": "Records subscription payment transactions.",
    "prompt_usage_logs": "Records system-prompt usage and model token counts.",
    "reward_transactions": "Records points awarded to students.",
    "roles": "Defines user roles and their descriptions.",
    "session_states": "Stores the current tutoring state for each AI session.",
    "show_me_breakdowns": "Stores worked examples and step-by-step explanations.",
    "speech_to_text_logs": "Records speech-to-text processing results.",
    "student_attempts": "Records student answers, scores, and attempt outcomes.",
    "student_avatar_items": "Associates students with unlocked avatar items.",
    "student_help_events": "Records help requests and strategies triggered during tutoring.",
    "student_inputs": "Stores typed, image, and audio inputs submitted by students.",
    "student_login_challenges": "Stores temporary QR and PIN login challenges for students.",
    "student_parent_links": "Associates students with parents or guardians.",
    "student_profiles": "Stores student-specific identity, class, and learning information.",
    "student_section_enrollments": "Records student section enrollment history.",
    "student_skill_features": "Stores calculated learning features for each student and skill.",
    "student_skill_predictions": "Stores predicted mastery and risk levels for student skills.",
    "student_teacher_links": "Associates students with teachers and class sections.",
    "subscription_plans": "Defines subscription products, prices, and access limits.",
    "subscription_students": "Associates subscriptions with students receiving access.",
    "subscriptions": "Stores purchased subscription access periods.",
    "system_prompts": "Stores versioned prompts used by the AI tutor.",
    "teacher_profiles": "Stores teacher-specific profile information.",
    "text_to_speech_logs": "Records generated speech audio for AI messages.",
    "user_auth": "Represents the temporary authentication-layer identifier.",
    "users": "Stores general user account information and assigned roles.",
    "vector_records": "Stores references to embeddings in semantic-memory storage.",
}


def split_top_level(text):
    parts, current, depth, quote = [], [], 0, None
    for ch in text:
        if quote:
            current.append(ch)
            if ch == quote:
                quote = None
        elif ch in "'\"":
            quote = ch
            current.append(ch)
        elif ch == "(":
            depth += 1
            current.append(ch)
        elif ch == ")":
            depth -= 1
            current.append(ch)
        elif ch == "," and depth == 0:
            parts.append("".join(current).strip())
            current = []
        else:
            current.append(ch)
    if "".join(current).strip():
        parts.append("".join(current).strip())
    return parts


def table_blocks(sql):
    pattern = re.compile(r"CREATE\s+TABLE\s+(?:public\.)?(\w+)\s*\(", re.I)
    for match in pattern.finditer(sql):
        start, depth, quote, i = match.end(), 1, None, match.end()
        while i < len(sql) and depth:
            ch = sql[i]
            if quote:
                if ch == quote:
                    quote = None
            elif ch in "'\"": quote = ch
            elif ch == "(": depth += 1
            elif ch == ")": depth -= 1
            i += 1
        yield match.group(1), sql[start:i - 1]


def parse_schema(sql):
    tables = {}
    for table, body in table_blocks(sql):
        primary_keys = []
        for item in split_top_level(body):
            cleaned = re.sub(r"--[^\n]*", "", item).strip()
            col = re.match(r"^(\w+)\s+", cleaned)
            if col and re.search(r"\bPRIMARY\s+KEY\b", cleaned, re.I):
                primary_keys.append(col.group(1))
            composite = re.search(r"PRIMARY\s+KEY\s*\(([^)]+)\)", cleaned, re.I)
            if composite and cleaned.upper().startswith(("PRIMARY KEY", "CONSTRAINT")):
                primary_keys.extend(x.strip() for x in composite.group(1).split(","))
        tables[table] = {"primary_keys": list(dict.fromkeys(primary_keys))}
    return tables


def parse_relationships(sql):
    relationships = []
    for alter in re.finditer(r"ALTER\s+TABLE\s+(?:public\.)?(\w+)\s+([\s\S]*?);", sql, re.I):
        child, body = alter.groups()
        for match in re.finditer(
            r"FOREIGN\s+KEY\s*\(([^)]+)\)\s*REFERENCES\s+(?:public\.)?(\w+)\s*\(([^)]+)\)",
            body,
            re.I,
        ):
            child_cols, parent, parent_cols = match.groups()
            child_cols = ", ".join(x.strip() for x in child_cols.split(","))
            parent_cols = ", ".join(x.strip() for x in parent_cols.split(","))
            relationships.append({
                "parent": parent,
                "child": child,
                "child_cols": child_cols,
                "parent_cols": parent_cols,
            })
    return sorted(relationships, key=lambda x: (x["parent"].lower(), x["child"].lower(), x["child_cols"].lower()))


def set_font(run):
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")


def set_cell_margins(cell, value=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def hide_borders(table):
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    props.append(borders)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    props.append(element)


def build_document(path, headers, rows, widths):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

    opening = doc.add_paragraph()
    opening.paragraph_format.space_after = Pt(0)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    hide_borders(table)
    repeat_header(table.rows[0])

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        set_font(paragraph.add_run(header))

    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(paragraph.add_run(value or "—"))
    doc.save(path)


sql = SQL_PATH.read_text(encoding="utf-8")
tables = parse_schema(sql)
relationships = parse_relationships(sql)

foreign_keys_by_table = {table: [] for table in tables}
for rel in relationships:
    foreign_keys_by_table.setdefault(rel["child"], []).append(
        f"{rel['child_cols']} → {rel['parent']}.{rel['parent_cols']}"
    )

schema_rows = []
for table in sorted(tables, key=str.lower):
    schema_rows.append([
        table,
        ", ".join(tables[table]["primary_keys"]) or "—",
        "; ".join(foreign_keys_by_table.get(table, [])) or "—",
        TABLE_DESCRIPTIONS.get(table, f"Stores records for {table.replace('_', ' ')}."),
    ])

relationship_rows = []
one_to_one = {
    ("users", "student_profiles", "user_id"),
    ("users", "billing_accounts", "user_id"),
    ("organizations", "billing_accounts", "organization_id"),
}
for rel in relationships:
    fk = f"{rel['child']}.{rel['child_cols']} → {rel['parent']}.{rel['parent_cols']}"
    relationship_type = (
        "One-to-one"
        if (rel["parent"], rel["child"], rel["child_cols"]) in one_to_one
        else "One-to-many"
    )
    relationship_rows.append([
        rel["parent"],
        rel["child"],
        relationship_type,
        fk,
        f"Connects {rel['child'].replace('_', ' ')} records to their related {rel['parent'].replace('_', ' ')} record.",
    ])

build_document(
    SCHEMA_OUTPUT,
    ["Table Name", "Primary Key", "Foreign Key(s)", "Description"],
    schema_rows,
    [1.15, 1.10, 1.75, 2.00],
)
build_document(
    REL_OUTPUT,
    ["Parent Table", "Child Table", "Relationship", "Foreign Key", "Description"],
    relationship_rows,
    [1.05, 1.10, 0.90, 1.40, 1.55],
)

print(f"Created schema document: {len(schema_rows)} tables")
print(f"Created relationship document: {len(relationship_rows)} foreign keys")
