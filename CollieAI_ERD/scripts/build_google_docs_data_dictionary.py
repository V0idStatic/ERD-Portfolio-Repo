import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "DATA_DICTIONARY.md"
OUTPUT = ROOT / "CollieAI_Data_Dictionary_Latest_Exact.docx"

CATEGORIES = {
    "Identity and Access": {
        "user_auth", "roles", "users", "teacher_profiles", "student_profiles",
        "parent_profiles", "student_parent_links", "student_teacher_links",
        "student_login_challenges",
    },
    "Organization and Classroom": {
        "organizations", "organization_members", "organization_student_enrollments",
        "class_sections", "student_section_enrollments",
    },
    "Subscription and Billing": {
        "subscription_plans", "billing_accounts", "subscriptions",
        "subscription_students", "payment_transactions",
    },
    "Learning Content": {
        "math_topics", "math_skills", "difficulty_levels", "math_problems",
    },
    "Multimodal Input": {
        "student_inputs", "file_assets", "ocr_logs", "speech_to_text_logs",
        "text_to_speech_logs",
    },
    "AI Tutoring and State": {
        "ai_sessions", "ai_messages", "session_states", "student_attempts",
        "attempt_steps", "dda_events", "show_me_breakdowns", "student_help_events",
    },
    "Guardrails and Prompt Safety": {
        "guardrail_checks", "system_prompts", "prompt_usage_logs",
    },
    "Machine Learning and Recommendations": {
        "student_skill_features", "student_skill_predictions", "ml_models",
        "learning_recommendations",
    },
    "Gamification and Rewards": {
        "avatars", "reward_transactions", "avatar_items", "student_avatar_items",
    },
    "Orchestration Logs and Notifications": {
        "n8n_workflow_logs", "api_request_logs", "notifications",
    },
    "Semantic Memory": {"vector_records"},
}


def category_for(table):
    for category, members in CATEGORIES.items():
        if table in members:
            return category
    return "Other"


def parse_markdown():
    text = SOURCE.read_text(encoding="utf-8")
    chunks = re.split(r"(?m)^## `([^`]+)`\s*$", text)
    tables = []
    for i in range(1, len(chunks), 2):
        name, body = chunks[i], chunks[i + 1]
        rows = []
        for line in body.splitlines():
            if not line.startswith("| `"):
                continue
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if len(cells) != 8:
                continue
            cells = [c.replace("\\|", "|").strip("`") for c in cells]
            rows.append(cells)
        tables.append((name, rows))
    return tables


def set_font(run, name="Times New Roman", size=9, italic=False, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.italic = italic
    run.font.bold = bold
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def set_cell_margins(cell, top=90, start=80, bottom=90, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "bottom", "left", "right", "insideV", "insideH"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def display_size(data_type, current_size):
    match = re.search(r"(?:VARCHAR|CHAR)\((\d+)\)", data_type, re.I)
    if match:
        return match.group(1)
    match = re.search(r"NUMERIC\((\d+)\s*,\s*(\d+)\)", data_type, re.I)
    if match:
        return f"{match.group(1)},{match.group(2)}"
    return "—"


def clean(value):
    return value.replace("`", "").replace("±", "+/-")


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

heading3 = doc.styles["Heading 3"]
heading3.font.name = "Times New Roman"
heading3.font.size = Pt(12)
heading3.font.bold = False
heading3.font.color.rgb = None
heading3._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
heading3._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
heading3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
heading3.paragraph_format.space_before = Pt(12)
heading3.paragraph_format.space_after = Pt(8)
heading3.paragraph_format.line_spacing = 1

headers = ["Field Name", "Data Type", "Data Format", "Field Size", "Description", "Example"]
widths = [1.0] * 6

for number, (table_name, rows) in enumerate(parse_markdown(), start=1):
    separator = doc.add_paragraph()
    separator.paragraph_format.space_after = Pt(0)
    if number > 1:
        separator.add_run().add_break(WD_BREAK.PAGE)

    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table)
    repeat_header(table.rows[0])

    for col_idx, header in enumerate(headers):
        cell = table.rows[0].cells[col_idx]
        cell.width = Inches(widths[col_idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=80, bottom=80)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        set_font(paragraph.add_run(header), size=11)

    for source in rows:
        field_name, data_type, data_format, field_size, _constraints, _required, desc, example = source
        output = [
            clean(field_name),
            clean(data_type).upper(),
            clean(data_format),
            clean(field_size),
            clean(desc),
            clean(example),
        ]
        cells = table.add_row().cells
        for col_idx, value in enumerate(output):
            cell = cells[col_idx]
            cell.width = Inches(widths[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx in (0, 4, 5) else WD_ALIGN_PARAGRAPH.CENTER
            set_font(paragraph.add_run(value), size=11)

doc.save(OUTPUT)
print(f"Created {OUTPUT.name} with {len(parse_markdown())} copy-ready tables")
