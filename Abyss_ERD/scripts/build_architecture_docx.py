from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "ABYSS_Architecture_Google_Docs.docx"
BLACK, HEADER, BORDER = "000000", "E7E7E7", "BFBFBF"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def cell_margins(cell, value=90):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    props = table._tbl.tblPr
    border_box = props.first_child_found_in("w:tblBorders")
    if border_box is None:
        border_box = OxmlElement("w:tblBorders")
        props.append(border_box)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = border_box.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            border_box.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), BORDER)


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    for row_index, row in enumerate(table.rows):
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for column_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell)
            if row_index == 0:
                shade(cell, HEADER)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.2)
                    run.font.color.rgb = RGBColor.from_string(BLACK)
                    run.bold = row_index == 0
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.62)
    section.left_margin = section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = document.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(10)
    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    document.add_paragraph("ABYSS System Architecture", style="Title")
    document.add_paragraph(
        "This architecture uses one React Native codebase for Android, iOS, and web; FastAPI for the REST API; Python workers for statistics and matching; Supabase for managed PostgreSQL and object storage; Vercel for the web build; and Render for the API and background workers."
    )

    document.add_heading("System Architecture", level=1)
    document.add_paragraph(
        "The mobile and web clients communicate with FastAPI through HTTPS REST requests. FastAPI validates authentication, permissions, game compatibility, schedules, and business rules before it changes the normalized entity tables. The clients do not write directly to PostgreSQL."
    )
    document.add_paragraph(
        "Live display endpoints use vw_ views, while dashboards and rankings use mv_ materialized views. Python workers process database events, calculate statistics and ratings, refresh materialized views, prepare historical features, and train future matching models separately from normal API requests."
    )
    add_table(document, ["Layer", "Component", "Communicates with", "Information exchanged"], [
        ("Users", "Players and coaches", "React Native mobile app", "Profiles, squad registration, matchmaking, invitations, schedules, and results"),
        ("Users", "Moderators and administrators", "React Native Web app", "Reports, verification, moderation, and dashboard operations"),
        ("Frontend", "Mobile and web applications", "FastAPI REST API on Render", "HTTPS requests and JSON responses"),
        ("Backend", "FastAPI", "Supabase PostgreSQL", "Validated transactions and display queries"),
        ("Backend", "FastAPI profile module", "Supabase Storage", "Signed uploads and stable media object keys"),
        ("Database", "Normalized entity tables", "vw_ display views", "Readable profile, roster, schedule, message, and free-agent data"),
        ("Database", "Match and rating history", "Python statistics worker", "Completed results, participants, statistic values, and rating changes"),
        ("Background", "Python statistics worker", "mv_ materialized views", "Win rates, squad statistics, leaderboard ranks, and dashboard totals"),
        ("Intelligence", "Feature builder", "Matching model", "Game, mode, region, availability, rating, win rate, experience, and roster features"),
        ("Intelligence", "Matching model", "FastAPI matching endpoint", "Ranked compatible squads, compatibility score, and matching reasons"),
    ], [1.0, 1.65, 1.85, 2.9])

    document.add_heading("Technology Stack", level=1)
    add_table(document, ["Layer", "Language and framework", "Purpose"], [
        ("Mobile and web frontend", "TypeScript, React Native, Expo, Expo Router", "Shared Android, iOS, and browser application"),
        ("Client data", "TanStack Query and Zustand", "REST caching, request state, authentication state, and local interface state"),
        ("REST backend", "Python, FastAPI, Pydantic", "Authentication, validation, authorization, and JSON endpoints"),
        ("Database access", "SQLAlchemy 2, Alembic, psycopg", "PostgreSQL transactions, mappings, pooling, and migrations"),
        ("Statistics", "Python, pandas, NumPy", "Squad totals, win rates, ratings, features, and historical analysis"),
        ("Match recommendation", "scikit-learn and joblib", "Train, evaluate, version, and load the compatibility model"),
        ("Database", "Supabase PostgreSQL", "Entity tables, vw_ display views, and mv_ calculated displays"),
        ("File storage", "Supabase Storage", "Avatars, banners, posts, result screenshots, and verification evidence"),
        ("Web deployment", "Vercel", "React Native Web production build and static assets"),
        ("API and workers", "Render Web Service, Worker, and Cron Job", "FastAPI, events, calculations, refreshes, and scheduled training"),
        ("Mobile delivery", "Expo Application Services and app stores", "Build and distribute Android and iOS applications"),
        ("Version control", "Git and GitHub", "Source management, collaboration, and deployment integration"),
    ], [1.7, 2.35, 3.35])

    document.add_heading("Usable Assets", level=1)
    add_table(document, ["Available asset", "How it is used"], [
        ("ABYSS.pdf requirements and interface references", "Source for roles, screens, squad profiles, invitations, results, moderation, and dashboard behavior"),
        ("abyss.sql", "PostgreSQL implementation of the domain model, constraints, relationships, views, and materialized views"),
        ("Interactive ERD in dist", "Visual reference for frontend, backend, database, and reporting implementation"),
        ("vw_ display objects", "Display-ready profile, roster, schedule, message, free-agent, and matchmaking data"),
        ("mv_ calculated objects", "Precalculated squad statistics, win rates, leaderboard ranks, and administrator totals"),
        ("Historical match records", "Training inputs from match participants, participant statistics, results, and rating changes"),
        ("Supabase Storage", "Managed storage for public profile media and protected evidence"),
        ("React Native and Expo libraries", "Reusable navigation, forms, images, secure local storage, responsive layouts, and mobile builds"),
        ("Python data libraries", "Reusable statistics, feature preparation, model training, and evaluation tools"),
    ], [2.55, 4.85])
    document.add_paragraph(
        "Logos, finalized brand files, production game catalogs, and a trained matching model should only be listed as usable assets after the team has created or obtained them."
    )

    document.add_heading("Main System Flow", level=1)
    for item in [
        "A coach registers a squad and its five players through the React Native application.",
        "FastAPI creates the squad, coach, player profiles, game profiles, and squad memberships in Supabase PostgreSQL.",
        "The coach requests compatible opponents from the matchmaking endpoint.",
        "FastAPI reads vw_matchmaking_squads and mv_squad_statistics, applies eligibility filters, and requests compatibility scores.",
        "The matching service returns ranked squads with a score and understandable matching reasons.",
        "The coach sends an invitation, and FastAPI validates coach authority, game compatibility, and schedule availability before saving it.",
        "After the scrim, the coach submits the result, evidence, participating players, and available player statistics.",
        "FastAPI saves the historical match records and finalizes an accepted result.",
        "A Python worker calculates squad totals, win rates, rating changes, and refreshed materialized views.",
        "Historical records are converted into features for scheduled matching-model training and evaluation.",
        "The frontend requests the updated profile, dashboard, schedule, or leaderboard and receives display-ready JSON data.",
    ]:
        document.add_paragraph(item, style="List Number")

    document.add_heading("Database Responsibility", level=1)
    add_table(document, ["Purpose", "ERD objects"], [
        ("Permanent identity", "users, player_profiles, player_profile_claims"),
        ("Per-game player data", "games, game_modes, game_roles, game_ranks, game_characters, player_game_profiles"),
        ("Squad history", "squads, squad_coaches, squad_memberships"),
        ("Match history", "scrim_invites, scrims, match_participants, scrim_results"),
        ("Flexible game statistics", "stat_definitions, participant_stat_values, rating_changes"),
        ("Live display data", "vw_ views"),
        ("Stored calculations", "mv_squad_statistics, mv_leaderboard, mv_admin_dashboard_metrics"),
        ("Reliable background processing", "event_outbox, event_receipts, request_receipts"),
    ], [2.2, 5.2])

    document.add_heading("Scalability and Matching Approach", level=1)
    document.add_paragraph(
        "The schema supports additional games without changing the main ERD because games, modes, roles, ranks, characters, and statistic definitions are configurable records. FastAPI and the Python workers can scale independently so calculations and model training do not delay ordinary REST requests."
    )
    document.add_paragraph(
        "The first matching version should use deterministic eligibility filters and a weighted score based on the same game and mode, region, availability overlap, squad rating difference, recent win rate, completed-match confidence, and roster completeness. Machine learning should be introduced after enough accepted match history exists. Training runs asynchronously, every model is versioned, and a new model is used only after its evaluation improves on the existing scoring baseline."
    )
    document.add_paragraph(
        "React Native Web is deployed to Vercel. Android and iOS builds are distributed as installed applications and call the same Render API. Render runs FastAPI separately from the worker and scheduled training process so API traffic is not blocked by statistics refreshes or training."
    )

    document.core_properties.title = "ABYSS System Architecture"
    document.core_properties.subject = "Architecture Usable Assets and Technology Stack"
    document.core_properties.author = "ABYSS Project Team"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
